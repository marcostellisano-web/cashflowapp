import unittest
from openpyxl import load_workbook
from app.models.budget import ParsedBudget, BudgetLineItem, BudgetDetailRow
from app.services.tax_credit_writer import write_tax_credit_excel, write_bible_excel


class FilingBibleTests(unittest.TestCase):
    def test_effective_rules_and_account_links(self):
        budget = ParsedBudget(
            source_filename='test.xlsx', total_budget=300,
            line_items=[BudgetLineItem(code='0201', description='Writer', total=300)],
            detail_rows=[BudgetDetailRow(account=code, description='Fee', subtotal=100, currency='CAD')
                         for code in ('02.01', '0201', '9999')])
        wb = load_workbook(write_tax_credit_excel(
            budget, 'Test',
            overrides={'0201': {'prov_labour_pct': 0, 'is_non_prov': False}},
            global_bible={'0201': (True, .5, .6, .1, .2, .3)},
            bible_descriptions={'0201': 'Custom writer'}))
        bible = wb['Bible']
        rules = {row[0]: row for row in bible.iter_rows(min_row=5, values_only=True)}
        self.assertEqual(rules['0201'][1:], ('Custom writer', None, None, .6, .1, .2, .3))
        self.assertIn('9999', rules)
        breakout = wb['Breakout Budget']
        detail = [row for row in breakout if row[0].value in ('02.01', '0201', '9999')]
        self.assertEqual(len(detail), 3)
        for row in detail:
            links = [cell for cell in row if isinstance(cell.value, str) and "'Bible'!" in cell.value]
            self.assertEqual(len(links), 6)
            for cell in links:
                self.assertIn('MATCH(SUBSTITUTE(SUBSTITUTE($A', cell.value)
                self.assertIn(',0)', cell.value)
        self.assertTrue(wb.calculation.fullCalcOnLoad)
        self.assertEqual(bible.freeze_panes, 'C5')

    def test_standalone_bible_still_exports(self):
        wb = load_workbook(write_bible_excel([{'account_code': '0201', 'is_non_prov': True}]))
        self.assertEqual(wb.sheetnames, ['Breakout Bible'])
        self.assertEqual(wb.active['C5'].value, 'OUT')

    def test_empty_budget(self):
        wb = load_workbook(write_tax_credit_excel(
            ParsedBudget(line_items=[], total_budget=0, source_filename='empty.xlsx'), 'Empty'))
        self.assertGreater(wb['Bible'].max_row, 5)


if __name__ == '__main__':
    unittest.main()
